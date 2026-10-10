import io

from docx import Document
from docx.shared import Pt

from preflight.graph.models import ProtocolGraph
from preflight.ingest.m11_map import load_toc


def _sort_key(m11_id: str) -> list[int]:
    return [int(x) if x.isdigit() else 0 for x in m11_id.split(".")]


def to_docx(graph: ProtocolGraph, sections: list[dict] | None = None) -> bytes:
    """Clean M11-ordered document. `sections` (m11_id, title, markdown) overrides the graph text, for drafts."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(graph.meta.title or "Clinical Trial Protocol", level=0)
    facts = [("NCT number", graph.meta.nct_id), ("Phase", graph.meta.phase), ("Condition", graph.meta.condition),
             ("Sponsor", graph.meta.sponsor), ("Masking", graph.meta.masking), ("Allocation", graph.meta.allocation)]
    for label, value in facts:
        if value:
            doc.add_paragraph(f"{label}: {value}")

    toc = {i["id"]: i["title"] for i in load_toc()}
    if sections is None:
        by_id: dict[str, list[str]] = {}
        for s in graph.sections:
            if s.m11_id:
                by_id.setdefault(s.m11_id, []).append(s.text)
        sections = [{"m11_id": k, "title": toc.get(k, k), "markdown": "\n".join(v)} for k, v in by_id.items()]
    for s in sorted(sections, key=lambda s: _sort_key(s["m11_id"])):
        doc.add_heading(f"{s['m11_id']} {s['title']}", level=min(s["m11_id"].count(".") + 1, 3))
        for para in s["markdown"].split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip().lstrip("#").strip())

    if graph.soa:
        doc.add_heading("Schedule of Activities", level=1)
        soa = graph.soa
        table = doc.add_table(rows=1 + len(soa.procedures), cols=1 + len(soa.visits))
        table.style = "Table Grid"
        table.cell(0, 0).text = "Procedure"
        for j, v in enumerate(soa.visits, 1):
            table.cell(0, j).text = v.name
        for i, p in enumerate(soa.procedures, 1):
            table.cell(i, 0).text = p.name
            for j, v in enumerate(soa.visits, 1):
                table.cell(i, j).text = "X" if soa.cells.get(v.id, {}).get(p.id) else ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
