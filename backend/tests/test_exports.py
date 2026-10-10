import csv
import io

import pymupdf
from docx import Document

from preflight.export import m11_docx, report, usdm
from tests.test_checks import graph, orr, soa

ISSUE = {"id": "iss_01", "check_id": "C1", "source": "rule", "category": "consistency", "severity": "critical",
         "title": "Endpoint at Week 24 has no visit", "status": "open",
         "where": {"section": "1.3 SoA", "page": 14, "bbox": None, "quote": "Week 24 imaging"},
         "why": {"explanation": "No visit.", "standard_ref": "ICH M11 1.3", "amendment_category": "visit_schedule"},
         "fix": {"suggestion": "Add a visit.", "patch": None}, "confidence": 1, "verified": True, "related": []}


def full_graph():
    ep, obj = orr([12])
    return graph(endpoints=[ep], objectives=[obj], soa=soa({"v2": {"p1": True}}))


def test_usdm_mapping():
    doc = usdm.to_usdm(full_graph())
    design = doc["study"]["versions"][0]["studyDesigns"][0]
    assert [e["name"] for e in design["encounters"]] == ["Screening", "Week 12", "Week 25"]
    assert design["scheduleTimelines"][0]["instances"][1]["activityIds"] == ["p1"]
    assert design["objectives"][0]["endpoints"][0]["purpose"] == "ORR"
    assert usdm.validate(doc) == []


def test_docx_has_soa_table():
    d = Document(io.BytesIO(m11_docx.to_docx(full_graph())))
    assert d.tables[0].cell(0, 2).text == "Week 12" and d.tables[0].cell(1, 2).text == "X"


def test_report_csv_and_pdf():
    rows = list(csv.DictReader(io.StringIO(report.to_csv([ISSUE]).decode())))
    assert rows[0]["id"] == "iss_01" and rows[0]["page"] == "14"
    run = {"run_id": "run_x", "mode": "full", "finished_at": None, "score": {"overall": 62}}
    pdf = pymupdf.open(stream=report.to_pdf("Trial", run, [ISSUE] * 40, ), filetype="pdf")
    assert pdf.page_count >= 2 and "iss_01" in pdf[0].get_text()
